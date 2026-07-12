from __future__ import annotations

from pathlib import Path

from app.core.logging import get_logger

logger = get_logger("ingestion.loader")

SUPPORTED_EXTENSIONS: set[str] = {".pdf", ".docx", ".txt", ".md", ".markdown"}


class UnsupportedFormatError(ValueError):
    def __init__(self, extension: str) -> None:
        super().__init__(f"Unsupported file format: {extension}")
        self.extension = extension


async def extract_text(file_path: str | Path) -> str:
    """Extract raw text from a file based on its extension."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".pdf":
        return await _extract_pdf(path)
    elif ext == ".docx":
        return await _extract_docx(path)
    elif ext in {".txt", ".md", ".markdown"}:
        return await _extract_text_file(path)
    else:
        raise UnsupportedFormatError(ext)


async def extract_text_from_bytes(
    content: bytes,
    filename: str,
) -> str:
    """Extract text from raw bytes (e.g. uploaded file)."""
    ext = Path(filename).suffix.lower()

    if ext == ".pdf":
        return await _extract_pdf_bytes(content)
    elif ext == ".docx":
        return await _extract_docx_bytes(content)
    elif ext in {".txt", ".md", ".markdown"}:
        return content.decode("utf-8", errors="replace")
    else:
        raise UnsupportedFormatError(ext)


async def _extract_pdf(path: Path) -> str:
    try:
        import pypdf
    except ImportError:
        raise ImportError("pypdf is required for PDF extraction: pip install pypdf")

    text_parts: list[str] = []
    reader = pypdf.PdfReader(str(path))
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    result = "\n\n".join(text_parts)
    logger.debug("extracted PDF text", extra={"path": str(path), "chars": len(result)})
    return result


async def _extract_pdf_bytes(content: bytes) -> str:
    import pypdf
    import io

    reader = pypdf.PdfReader(io.BytesIO(content))
    text_parts: list[str] = []
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text_parts.append(page_text)
    return "\n\n".join(text_parts)


async def _extract_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError:
        raise ImportError(
            "python-docx is required for DOCX extraction: pip install python-docx"
        )

    doc = Document(str(path))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    result = "\n".join(paragraphs)
    logger.debug("extracted DOCX text", extra={"path": str(path), "chars": len(result)})
    return result


async def _extract_docx_bytes(content: bytes) -> str:
    from docx import Document
    import io

    doc = Document(io.BytesIO(content))
    paragraphs = [p.text for p in doc.paragraphs if p.text]
    return "\n".join(paragraphs)


async def _extract_text_file(path: Path) -> str:
    text = path.read_text("utf-8", errors="replace")
    logger.debug("extracted text file", extra={"path": str(path), "chars": len(text)})
    return text
