from __future__ import annotations

import json
import uuid
from datetime import datetime, timezone
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.agents.state import ResearchState, make_initial_state
from app.agents.paper_author_agent import PaperAuthorAgent
from app.agents.proposal_agent import ProposalAgent
from app.agents.quality_review_agent import QualityReviewAgent
from app.agents.citation_validator_agent import CitationValidatorAgent
from app.agents.evidence_validator_agent import EvidenceValidatorAgent
from app.models.paper import (
    Proposal, Paper, PaperSection, PaperRevision, PaperCitation,
    PaperExport, PaperMetrics, PaperStatus, SectionStatus,
    PaperOperation,
)
from app.models.research_project import ResearchProject
from app.models.user import User
from app.schemas.paper import (
    ProposalCreate, ProposalResponse,
    PaperGenerateRequest, PaperResponse, PaperSectionResponse,
    PaperExportResponse,
    SectionRewriteRequest, SectionRewriteResponse,
    QualityReviewResponse, CitationValidationResponse,
    EvidenceValidationResponse, PaperListResponse,
)
from app.db.session import get_async_session
from app.core.logging import get_logger
from app.services.auth_service import get_current_user
from app.core.config import get_settings
from app.llm.factory import get_llm_provider
from app.llm.provider import LLMProvider

logger = get_logger("api.papers")
router = APIRouter(prefix="/papers", tags=["papers"])


async def _get_project(user: User, project_id: str, session: AsyncSession) -> ResearchProject:
    result = await session.execute(
        select(ResearchProject).where(
            ResearchProject.id == uuid.UUID(project_id),
            ResearchProject.created_by == user.id,
        )
    )
    project = result.scalar_one_or_none()
    if project is None:
        raise HTTPException(status_code=404, detail="Project not found")
    return project


async def _get_proposal(user: User, proposal_id: str, session: AsyncSession) -> Proposal:
    result = await session.execute(
        select(Proposal).where(Proposal.id == uuid.UUID(proposal_id))
    )
    proposal = result.scalar_one_or_none()
    if proposal is None:
        raise HTTPException(status_code=404, detail="Proposal not found")
    return proposal


async def _get_paper(user: User, paper_id: str, session: AsyncSession) -> Paper:
    result = await session.execute(
        select(Paper).where(Paper.id == uuid.UUID(paper_id))
    )
    paper = result.scalar_one_or_none()
    if paper is None:
        raise HTTPException(status_code=404, detail="Paper not found")
    return paper


def _get_llm() -> LLMProvider:
    settings = get_settings()
    return get_llm_provider(settings)


def _build_state_from_proposal(proposal: Proposal) -> ResearchState:
    state = make_initial_state(query=proposal.proposed_title, project_id=str(proposal.project_id))
    state["objective"] = proposal.problem_statement
    state["domain"] = proposal.domain or ""
    state["keywords"] = ", ".join(proposal.keywords) if proposal.keywords else ""
    state["selected_gap"] = {"description": proposal.problem_statement, "id": proposal.gap_id}
    state["proposal"] = {
        "proposed_title": proposal.proposed_title,
        "problem_statement": proposal.problem_statement,
        "motivation": proposal.motivation,
        "research_questions": proposal.research_questions,
        "hypothesis": proposal.hypothesis,
        "objectives": proposal.objectives,
        "expected_contributions": proposal.expected_contributions,
        "proposed_methodology": proposal.proposed_methodology,
        "evaluation_strategy": proposal.evaluation_strategy,
        "future_scope": proposal.future_scope,
        "keywords": proposal.keywords,
    }
    if proposal.base_paper_analysis:
        state["base_paper_context"] = json.dumps(proposal.base_paper_analysis, indent=2)
    return state


@router.post("/proposal", response_model=ProposalResponse, status_code=201)
async def create_proposal(
    request: ProposalCreate,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    _ = await _get_project(current_user, request.project_id, session)
    llm = _get_llm()
    agent = ProposalAgent(llm_provider=llm)

    state = make_initial_state(
        query=request.objective or request.domain,
        project_id=request.project_id,
        objective=request.objective,
    )
    state["selected_gap"] = {"description": request.gap_id or request.objective, "id": request.gap_id or ""}
    state["domain"] = request.domain
    state["keywords"] = request.keywords
    state["methodology_preference"] = request.methodology_preference

    base_paper_analysis = None
    if request.base_paper_doi or request.base_paper_url:
        analysis_result = await _analyze_base_paper(llm, request.base_paper_doi, request.base_paper_url)
        state["base_paper_context"] = json.dumps(analysis_result, indent=2)
        base_paper_analysis = analysis_result

    state = await agent.arun(state)
    proposal_data = state.get("proposal", {})

    db_proposal = Proposal(
        project_id=uuid.UUID(request.project_id),
        gap_id=request.gap_id,
        proposed_title=proposal_data.get("proposed_title", "Research Proposal"),
        problem_statement=proposal_data.get("problem_statement", ""),
        motivation=proposal_data.get("motivation", ""),
        research_questions=proposal_data.get("research_questions", []),
        hypothesis=proposal_data.get("hypothesis", ""),
        objectives=proposal_data.get("objectives", []),
        expected_contributions=proposal_data.get("expected_contributions", []),
        proposed_methodology=proposal_data.get("proposed_methodology", ""),
        evaluation_strategy=proposal_data.get("evaluation_strategy", ""),
        future_scope=proposal_data.get("future_scope", ""),
        keywords=proposal_data.get("keywords", []),
        domain=request.domain,
        base_paper_analysis=base_paper_analysis,
    )
    session.add(db_proposal)
    await session.commit()
    await session.refresh(db_proposal)

    logger.info("proposal created", extra={"proposal_id": str(db_proposal.id), "title": db_proposal.proposed_title})
    return db_proposal


async def _analyze_base_paper(llm, doi: str | None, url: str | None) -> dict[str, Any]:
    from app.agents.paper_authoring_prompts import (
        PAPER_AUTHOR_SYSTEM_PROMPT,
    )
    source = doi or url or ""
    prompt = f"""Analyze the following base paper and extract key information. Return JSON only.
Source: {source}

Extract: abstract, keywords, methodology, dataset_description, experiments, limitations, future_work, references (list of strings).

If you cannot access the full paper, extract what you can from the metadata and indicate uncertainty.

Return ONLY valid JSON with exactly these keys: abstract, keywords, methodology, dataset_description, experiments, limitations, future_work, references"""
    try:
        response = await llm.generate(prompt=prompt, system_prompt=PAPER_AUTHOR_SYSTEM_PROMPT)
        raw = response.content.strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1]
            raw = raw.rsplit("\n```", 1)[0]
        return json.loads(raw)
    except Exception as exc:
        logger.warning("base paper analysis failed", extra={"error": str(exc), "source": source})
        return {
            "abstract": None, "keywords": [], "methodology": None,
            "dataset_description": None, "experiments": None,
            "limitations": None, "future_work": None, "references": [],
            "originality_note": "Analysis incomplete — verify against original source.",
        }


@router.get("/proposal/{proposal_id}", response_model=ProposalResponse)
async def get_proposal(
    proposal_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    proposal = await _get_proposal(current_user, proposal_id, session)
    return proposal


@router.get("/proposals/{project_id}", response_model=list[ProposalResponse])
async def list_proposals(
    project_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    await _get_project(current_user, project_id, session)
    result = await session.execute(
        select(Proposal).where(Proposal.project_id == uuid.UUID(project_id))
    )
    return result.scalars().all()


@router.post("/generate", response_model=PaperResponse, status_code=201)
async def generate_paper(
    request: PaperGenerateRequest,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    await _get_project(current_user, request.project_id, session)
    proposal = await _get_proposal(current_user, request.proposal_id, session)

    llm = _get_llm()
    agent = PaperAuthorAgent(llm_provider=llm)
    state = _build_state_from_proposal(proposal)

    state = await agent.arun(state)
    paper_data = state.get("paper_draft", {})

    db_paper = Paper(
        project_id=uuid.UUID(request.project_id),
        proposal_id=proposal.id,
        title=paper_data.get("title", proposal.proposed_title),
        abstract=paper_data.get("abstract", ""),
        keywords=paper_data.get("keywords", []),
        status=PaperStatus.GENERATED,
        ieee_content=json.dumps(paper_data, indent=2),
    )
    session.add(db_paper)
    await session.flush()

    sections_data = paper_data.get("sections", [])
    for sec in sections_data:
        content = sec.get("content", "")
        db_section = PaperSection(
            paper_id=db_paper.id,
            section_number=sec.get("section_number", 1),
            section_title=sec.get("section_title", ""),
            content=content,
            word_count=len(content.split()),
            status=SectionStatus.DRAFT,
        )
        session.add(db_section)

    refs_data = paper_data.get("references", [])
    for ref in refs_data:
        db_citation = PaperCitation(
            paper_id=db_paper.id,
            citation_key=ref.get("citation_key", ""),
            authors=ref.get("authors", ""),
            title=ref.get("title", ""),
            year=ref.get("year"),
            journal=ref.get("journal"),
            doi=ref.get("doi"),
            url=ref.get("url"),
        )
        session.add(db_citation)

    db_metrics = PaperMetrics(paper_id=db_paper.id)
    session.add(db_metrics)

    await session.commit()
    await session.refresh(db_paper)

    logger.info("paper generated", extra={
        "paper_id": str(db_paper.id),
        "title": db_paper.title,
        "sections": len(sections_data),
        "references": len(refs_data),
    })
    return db_paper


@router.get("/{paper_id}", response_model=PaperResponse)
async def get_paper(
    paper_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)
    return paper


@router.put("/{paper_id}", response_model=PaperResponse)
async def update_paper(
    paper_id: str,
    title: str | None = None,
    abstract: str | None = None,
    keywords: list[str] | None = None,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)
    if title is not None:
        paper.title = title
    if abstract is not None:
        paper.abstract = abstract
    if keywords is not None:
        paper.keywords = keywords
    paper.updated_at = datetime.now(timezone.utc)
    await session.commit()
    await session.refresh(paper)
    return paper


@router.get("/project/{project_id}", response_model=PaperListResponse)
async def list_papers(
    project_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    await _get_project(current_user, project_id, session)
    result = await session.execute(
        select(Paper).where(Paper.project_id == uuid.UUID(project_id))
    )
    papers = result.scalars().all()
    return PaperListResponse(papers=papers, total=len(papers))


@router.post("/{paper_id}/sections/{section_id}/rewrite", response_model=SectionRewriteResponse)
async def rewrite_section(
    paper_id: str,
    section_id: str,
    request: SectionRewriteRequest,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)
    section = await session.get(PaperSection, uuid.UUID(section_id))
    if section is None or section.paper_id != paper.id:
        raise HTTPException(status_code=404, detail="Section not found")

    llm = _get_llm()
    agent = PaperAuthorAgent(llm_provider=llm)
    old_content = section.content

    operation_map = {
        "rewrite": agent.rewrite_section,
        "expand": agent.expand_section,
        "condense": agent.condense_section,
        "improve_tone": agent.improve_tone,
        "add_citations": lambda t, c: agent.add_citations(t, c, _get_citations_text(paper)),
        "improve_depth": agent.improve_depth,
    }

    op_func = operation_map.get(request.operation)
    if op_func is None:
        raise HTTPException(status_code=400, detail=f"Unknown operation: {request.operation}")

    new_content = await op_func(section.section_title, old_content)
    if new_content is None:
        new_content = old_content

    revision = PaperRevision(
        paper_id=paper.id,
        section_id=section.id,
        section_title=section.section_title,
        operation=PaperOperation(request.operation),
        old_content=old_content,
        new_content=new_content,
    )
    session.add(revision)

    section.content = new_content
    section.word_count = len(new_content.split())
    section.updated_at = datetime.now(timezone.utc)
    await session.commit()

    return SectionRewriteResponse(
        section_id=section.id,
        new_content=new_content,
        revision_id=revision.id,
        operation=request.operation,
    )


def _get_citations_text(paper: Paper) -> str:
    return "\n".join(
        f"{c.citation_key}: {c.authors} ({c.year}). {c.title}. {c.doi or c.url or ''}"
        for c in paper.citations
    ) if paper.citations else "No citations available."


@router.post("/{paper_id}/quality-review", response_model=QualityReviewResponse)
async def quality_review(
    paper_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)

    llm = _get_llm()
    agent = QualityReviewAgent(llm_provider=llm)
    state = make_initial_state(query=paper.title, project_id=str(paper.project_id))

    sections = paper.sections
    citations = paper.citations
    paper_draft = {
        "title": paper.title,
        "abstract": paper.abstract,
        "sections": [
            {"section_number": s.section_number, "section_title": s.section_title, "content": s.content}
            for s in sections
        ],
        "references": [
            {"citation_key": c.citation_key, "authors": c.authors, "title": c.title, "year": c.year, "doi": c.doi}
            for c in citations
        ],
    }
    state["paper_draft"] = paper_draft
    state = await agent.arun(state)
    review = state.get("quality_review", {})

    if paper.metrics:
        metrics = paper.metrics
    else:
        metrics = PaperMetrics(paper_id=paper.id)
        session.add(metrics)

    metrics.novelty_score = review.get("novelty_score", 0)
    metrics.citation_coverage = review.get("citation_coverage", 0)
    metrics.evidence_strength = review.get("evidence_strength", 0)
    metrics.methodology_quality = review.get("methodology_quality", 0)
    metrics.writing_quality = review.get("writing_quality", 0)
    metrics.logical_consistency = review.get("logical_consistency", 0)
    metrics.academic_tone = review.get("academic_tone", 0)
    metrics.section_completeness = review.get("section_completeness", 0)
    metrics.composite_score = review.get("composite_score", 0)
    metrics.details = review
    metrics.suggestions = review.get("suggestions", [])
    await session.commit()

    return QualityReviewResponse(
        novelty_score=metrics.novelty_score,
        citation_coverage=metrics.citation_coverage,
        evidence_strength=metrics.evidence_strength,
        methodology_quality=metrics.methodology_quality,
        writing_quality=metrics.writing_quality,
        logical_consistency=metrics.logical_consistency,
        academic_tone=metrics.academic_tone,
        section_completeness=metrics.section_completeness,
        composite_score=metrics.composite_score,
        suggestions=metrics.suggestions or [],
        strengths=review.get("strengths", []),
        weaknesses=review.get("weaknesses", []),
    )


@router.post("/{paper_id}/validate-citations", response_model=CitationValidationResponse)
async def validate_citations(
    paper_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)

    llm = _get_llm()
    agent = CitationValidatorAgent(llm_provider=llm)
    state = make_initial_state(query=paper.title, project_id=str(paper.project_id))

    refs = [
        {"citation_key": c.citation_key, "authors": c.authors, "title": c.title,
         "year": c.year, "journal": c.journal, "doi": c.doi, "url": c.url}
        for c in paper.citations
    ]
    state["paper_draft"] = {"references": refs}
    state = await agent.arun(state)
    validation = state.get("citation_validation", {})

    for c_data in validation.get("citations", []):
        key = c_data.get("citation_key", "")
        for db_c in paper.citations:
            if db_c.citation_key == key:
                db_c.verified = not c_data.get("is_fabricated", False) and not c_data.get("is_duplicate", False)
                db_c.ieee_format = c_data.get("ieee_format", "")
                db_c.verification_errors = c_data.get("verification_notes", [])
    await session.commit()

    return CitationValidationResponse(
        citations=validation.get("citations", []),
        total=validation.get("total", 0),
        duplicates=validation.get("duplicates", []),
        verified_count=validation.get("verified_count", 0),
        has_issues=validation.get("has_issues", False),
    )


@router.post("/{paper_id}/validate-evidence", response_model=EvidenceValidationResponse)
async def validate_evidence(
    paper_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)

    llm = _get_llm()
    agent = EvidenceValidatorAgent(llm_provider=llm)
    state = make_initial_state(query=paper.title, project_id=str(paper.project_id))

    sections = paper.sections
    state["paper_draft"] = {
        "sections": [
            {"section_number": s.section_number, "section_title": s.section_title, "content": s.content}
            for s in sections
        ],
    }
    state = await agent.arun(state)
    validation = state.get("evidence_validation", {})

    for sec_data in validation.get("sections", []):
        sec_num = sec_data.get("section_number")
        for db_sec in sections:
            if db_sec.section_number == sec_num:
                db_sec.evidence_classifications = sec_data

    coverage = validation.get("coverage", {})
    return EvidenceValidationResponse(
        sections=validation.get("sections", []),
        coverage=coverage,
        overall_supported_ratio=validation.get("overall_supported_ratio", 0),
        status=validation.get("status", "complete"),
    )


@router.get("/{paper_id}/download/{fmt}")
async def download_paper(
    paper_id: str,
    fmt: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)

    if fmt not in ("pdf", "docx"):
        raise HTTPException(status_code=400, detail="Unsupported format. Use 'pdf' or 'docx'.")

    existing = await session.execute(
        select(PaperExport).where(
            PaperExport.paper_id == paper.id,
            PaperExport.format == fmt,
        ).order_by(PaperExport.generated_at.desc())
    )
    export = existing.scalar_one_or_none()

    if export and export.file_path:
        from fastapi.responses import FileResponse
        import os
        if os.path.exists(export.file_path):
            return FileResponse(
                export.file_path,
                media_type="application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                filename=f"{paper.title[:50]}.{fmt}",
            )

    ieee_md = _render_ieee_markdown(paper)
    export_path = await _generate_export(paper, fmt, ieee_md, session)

    export_record = PaperExport(
        paper_id=paper.id,
        format=fmt,
        file_path=export_path,
    )
    session.add(export_record)
    await session.commit()

    from fastapi.responses import FileResponse
    return FileResponse(
        export_path,
        media_type="application/pdf" if fmt == "pdf" else "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=f"{paper.title[:50]}.{fmt}",
    )


def _render_ieee_markdown(paper: Paper) -> str:
    lines = []
    lines.append(f"# {paper.title}")
    lines.append("")
    lines.append(f"**Authors:** {paper.authors}")
    lines.append("")
    lines.append("## Abstract")
    lines.append("")
    lines.append(paper.abstract)
    lines.append("")
    lines.append(f"**Keywords:** {', '.join(paper.keywords)}")
    lines.append("")
    lines.append("---")
    lines.append("")

    for section in sorted(paper.sections, key=lambda s: s.section_number):
        lines.append(f"## {section.section_number}. {section.section_title}")
        lines.append("")
        lines.append(section.content)
        lines.append("")

    if paper.citations:
        lines.append("## References")
        lines.append("")
        for c in paper.citations:
            ieee = c.ieee_format or f"{c.authors} \"{c.title},\" {c.year}."
            lines.append(f"[{c.citation_key}] {ieee}")
        lines.append("")

    if paper.metrics:
        lines.append("---")
        lines.append(f"**Quality Score:** {paper.metrics.composite_score:.1f}/100")
        lines.append("")

    return "\n".join(lines)


async def _generate_export(paper: Paper, fmt: str, ieee_md: str, session: AsyncSession) -> str:
    import os
    export_dir = os.path.join("exports", str(paper.id))
    os.makedirs(export_dir, exist_ok=True)

    md_path = os.path.join(export_dir, "paper.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(ieee_md)

    if fmt == "pdf":
        pdf_path = os.path.join(export_dir, "paper.pdf")
        try:
            import markdown
            from weasyprint import HTML
            html = markdown.markdown(ieee_md, extensions=["tables", "fenced_code"])
            styled = f"""<html><head><style>
body {{ font-family: 'Times New Roman', serif; font-size: 12pt; line-height: 1.6; margin: 1in; }}
h1 {{ font-size: 18pt; text-align: center; }}
h2 {{ font-size: 14pt; margin-top: 1em; }}
table {{ border-collapse: collapse; width: 100%; }}
td, th {{ border: 1px solid #333; padding: 8px; }}
code {{ background: #f4f4f4; padding: 2px 4px; }}
</style></head><body>{html}</body></html>"""
            HTML(string=styled).write_pdf(pdf_path)
            return pdf_path
        except ImportError:
            fallback_path = os.path.join(export_dir, "paper.md")
            return fallback_path

    elif fmt == "docx":
        docx_path = os.path.join(export_dir, "paper.docx")
        try:
            from docx import Document
            from docx.shared import Pt
            doc = Document()
            style = doc.styles["Normal"]
            style.font.name = "Times New Roman"
            style.font.size = Pt(12)
            doc.add_heading(paper.title, 0)
            doc.add_paragraph(f"Authors: {paper.authors}")
            doc.add_heading("Abstract", 1)
            doc.add_paragraph(paper.abstract)
            if paper.keywords:
                doc.add_paragraph(f"Keywords: {', '.join(paper.keywords)}")
            for section in sorted(paper.sections, key=lambda s: s.section_number):
                doc.add_heading(f"{section.section_number}. {section.section_title}", 2)
                for para in section.content.split("\n\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())
            if paper.citations:
                doc.add_heading("References", 1)
                for c in paper.citations:
                    ieee = c.ieee_format or f"{c.authors}, \"{c.title},\" {c.year}."
                    doc.add_paragraph(f"[{c.citation_key}] {ieee}")
            doc.save(docx_path)
            return docx_path
        except ImportError:
            return md_path

    return md_path


@router.put("/{paper_id}/section/{section_id}", response_model=PaperSectionResponse)
async def update_section_content(
    paper_id: str,
    section_id: str,
    content: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)
    section = await session.get(PaperSection, uuid.UUID(section_id))
    if section is None or section.paper_id != paper.id:
        raise HTTPException(status_code=404, detail="Section not found")
    section.content = content
    section.word_count = len(content.split())
    section.updated_at = datetime.now(timezone.utc)
    await session.commit()
    await session.refresh(section)
    return section


@router.get("/{paper_id}/exports", response_model=list[PaperExportResponse])
async def list_exports(
    paper_id: str,
    current_user: User = Depends(get_current_user),
    session: AsyncSession = Depends(get_async_session),
) -> Any:
    paper = await _get_paper(current_user, paper_id, session)
    return paper.exports
